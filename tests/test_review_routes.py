import io
from datetime import date, timedelta

import pytest
from bs4 import BeautifulSoup
from factories import (
    DRAFT_FOR_ROUTES,
    PROPOSAL,
    FakeModel,
    add_converted_source,
    cited,
    offer,
    question,
    session_key_of,
)
from openpyxl import load_workbook

from build_a_hia.services.generation import GenerationProcessor

EMPTY_DRAFT = {"offers": [], "questions": [], "issues": []}


class RoutingModel(FakeModel):
    """Answers per sub-category, so the order of generation jobs does not matter."""

    def __init__(self, by_subcategory: dict[str, dict]) -> None:
        super().__init__()
        self.by_subcategory = by_subcategory

    def parse(self, messages, schema):
        task = messages[-1]["content"]
        name = next(name for name in self.by_subcategory if f"Sub-category: {name}" in task)
        self.responses.insert(0, self.by_subcategory[name])
        return super().parse(messages, schema)


@pytest.fixture
def run_jobs(services):
    def run(*responses, by_subcategory: dict[str, dict] | None = None) -> None:
        model = RoutingModel(by_subcategory) if by_subcategory else FakeModel(*responses)
        processor = GenerationProcessor(
            table=services.table,
            blobs=services.blobs,
            queue=services.storage.generate_queue,
            sessions=services.sessions,
            structures=services.structures,
            contents=services.contents,
            settings=services.settings,
            model_factory=lambda: model,
        )
        while processor.process_next():
            pass

    return run


@pytest.fixture
def generated(ready_client, run_jobs, services):
    """Approved structure with generated content: Night shelter (1 offer) and Cash (empty)."""
    ready_client.post("/structure/propose", data={"instructions": ""})
    run_jobs(PROPOSAL)
    ready_client.post("/structure/approve", data={"version": 1})
    ready_client.post("/content/generate")
    run_jobs(by_subcategory={"Night shelter": DRAFT_FOR_ROUTES, "Cash assistance": EMPTY_DRAFT})
    key = session_key_of(ready_client)
    structure = services.structures.approved(key)[1]
    night, cash = (sub for _category, sub in structure.subcategories())
    return ready_client, key, night.key, cash.key


def _state(services, key, sub_key):
    return services.contents.state(key, sub_key)


def _offer_form(blob: str, **values: str) -> dict[str, str]:
    data = {
        "blob": blob,
        "name": "Night shelter",
        "description": "Free beds for everyone.",
        "phone_numbers": "",
        "emails": "",
        "web_urls": "",
        "address": "",
        "open_weekdays": "",
        "open_weekend": "",
        "need_to_know": "",
        "more_info": "",
        "chapter": "",
    }
    data.update(values)
    return data


def _approve_all(client, services, key, night, cash):
    client.post(
        f"/content/{cash}/empty",
        data={"blob": _state(services, key, cash).blob, "decision": "drop"},
    )
    client.post(
        f"/content/{night}/approve",
        data={"blob": _state(services, key, night).blob, "disclosure_reviewed": "y"},
    )


def test_pages_do_not_describe_generated_content_as_hidden(generated, services):
    client, key, night, cash = generated

    content_page = client.get(f"/content/{night}")
    assert content_page.status_code == 200
    assert b"visibility Hide" not in content_page.data

    _approve_all(client, services, key, night, cash)
    client.post("/download")
    download_page = client.get("/download")
    assert download_page.status_code == 200
    assert b"Latest snapshot" in download_page.data
    assert b"visibility" not in download_page.data
    assert b"Hidden rows" not in download_page.data
    assert b"hidden scaffolding row" not in download_page.data
    assert b"copy rows 2 and below" in download_page.data
    assert b"Paste special" in download_page.data
    assert b"Values only" in download_page.data
    assert b"column B" in download_page.data


def test_reviewer_edits_are_flagged_and_withdraw_approval(generated, services):
    client, key, night, _cash = generated
    state = _state(services, key, night)
    client.post(f"/content/{night}/approve", data={"blob": state.blob, "disclosure_reviewed": "y"})
    assert _state(services, key, night).approved
    content = services.contents.result(key, _state(services, key, night))
    offer_id = content.offers[0].id

    response = client.post(
        f"/content/{night}/offers/{offer_id}", data=_offer_form(_state(services, key, night).blob)
    )

    assert response.status_code == 302
    state = _state(services, key, night)
    assert state.edited and not state.approved
    edited = services.contents.result(key, state).offers[0]
    assert edited.description.edited
    assert edited.description.evidence
    page = client.get(f"/content/{night}")
    assert b"Edited by reviewer" in page.data


@pytest.mark.parametrize("item_type", ["offers", "questions"])
def test_content_autosave_returns_updated_version_and_rejects_stale_edits(
    generated, services, item_type
):
    client, key, night, _cash = generated
    state = _state(services, key, night)
    content = services.contents.result(key, state)
    if item_type == "offers":
        item_id = content.offers[0].id
        data = _offer_form(state.blob, description="Updated description.")
    else:
        item_id = content.questions[0].ref
        data = {
            "blob": state.blob,
            "question": "Who can stay?",
            "answer": "Everyone.",
            "parent": "",
        }
    endpoint = f"/content/{night}/{item_type}/{item_id}"
    headers = {"Accept": "application/json"}

    response = client.post(endpoint, data=data, headers=headers)

    assert response.status_code == 200
    assert response.json == {"saved": True, "blob": _state(services, key, night).blob}
    assert response.json["blob"] != state.blob
    assert _state(services, key, night).edited
    assert not _state(services, key, night).approved
    stale = client.post(endpoint, data=data, headers=headers)
    assert stale.status_code == 409
    assert stale.json["saved"] is False
    data["blob"] = response.json["blob"]
    assert client.post(endpoint, data=data, headers=headers).status_code == 200
    data["blob"] = _state(services, key, night).blob
    data["name" if item_type == "offers" else "question"] = ""
    invalid = client.post(endpoint, data=data, headers=headers)
    assert invalid.status_code == 400
    assert invalid.json["saved"] is False


def test_stale_edits_are_rejected(generated, services):
    client, key, night, _cash = generated
    content = services.contents.result(key, _state(services, key, night))
    old_blob = _state(services, key, night).blob
    client.post(f"/content/{night}/offers/new", data=_offer_form(old_blob, name="Legal aid"))

    response = client.post(
        f"/content/{night}/offers/{content.offers[0].id}/remove",
        data={"blob": old_blob},
        follow_redirects=True,
    )

    assert b"changed in the meantime" in response.data
    assert len(services.contents.result(key, _state(services, key, night)).offers) == 2


def test_approval_needs_the_current_version(generated, services):
    client, key, night, _cash = generated
    old_blob = _state(services, key, night).blob
    client.post(f"/content/{night}/offers/new", data=_offer_form(old_blob, name="Legal aid"))

    client.post(f"/content/{night}/approve", data={"blob": old_blob})
    assert not _state(services, key, night).approved

    client.post(f"/content/{night}/approve", data={"blob": _state(services, key, night).blob})
    assert _state(services, key, night).approved


def test_content_page_puts_gaps_after_content_and_approve_next_to_regenerate(generated):
    client, _key, night, _cash = generated

    page = client.get(f"/content/{night}").data.decode()

    assert page.index('id="offers"') < page.index('id="questions"') < page.index('id="gaps"')
    assert page.index('id="gaps"') < page.index("Approve content")
    assert page.index("Approve content") < page.index("Regenerate this sub-category")
    assert "disclosure_reviewed" not in page
    assert "data-content-autosave" in page
    assert "Edit offer" not in page
    assert "Edit question" not in page
    assert "Save offer" not in page
    assert "Save question" not in page
    for label in ("Add offer", "Remove offer"):
        assert f'class="button secondary">{label}</button>' in page
    offer_card = page.split('class="offer"', 1)[1].split("</article>", 1)[0]
    assert 'class="content-remove"' in offer_card
    assert "Remove question" in page
    assert 'name="phone_numbers"' in page
    assert 'name="answer"' in page
    assert "Follow-up to" not in page


def test_questions_can_be_added_and_nested(generated, services):
    client, key, night, _cash = generated
    page = client.get(f"/content/{night}").data.decode()
    subquestion_form = page.split('id="add-sub-question-q1"', 1)[1].split("</details>", 1)[0]
    assert "Add a sub-question" in subquestion_form
    assert 'type="hidden" name="parent" value="q1"' in subquestion_form
    assert 'name="blob"' in subquestion_form
    assert "<select" not in subquestion_form
    blob = _state(services, key, night).blob
    client.post(
        f"/content/{night}/questions/new",
        data={"blob": blob, "question": "How long?", "answer": "One week.", "parent": "q1"},
    )

    content = services.contents.result(key, _state(services, key, night))
    assert [(q.question, q.parent_ref) for q in content.questions] == [
        ("Who can stay?", ""),
        ("How long?", "q1"),
    ]
    page = client.get(f"/content/{night}").data.decode()
    assert 'class="question sub-question"' in page
    assert 'id="add-sub-question-q1"' in page
    assert f'id="add-sub-question-{content.questions[1].ref}"' not in page
    document = BeautifulSoup(page, "html.parser")
    parent_card = document.find("article", id="question-q1")
    child_card = document.find("article", id=f"question-{content.questions[1].ref}")
    assert parent_card is not None and child_card is not None
    assert child_card.find_parent("article") is parent_card
    assert len(document.select("#questions > article.question")) == 1
    assert len(document.select("article.sub-question")) == 1
    assert not document.select("form form")
    for card, label in ((parent_card, "Remove question"), (child_card, "Remove sub-question")):
        remove_form = card.find("form", class_="content-remove", recursive=False)
        assert remove_form is not None
        button = remove_form.find("button", class_="secondary")
        assert button is not None and button.get_text(strip=True) == label
    for card, parent in ((parent_card, ""), (child_card, "q1")):
        edit_form = card.find("form", attrs={"data-content-autosave": True}, recursive=False)
        assert edit_form is not None
        parent_field = edit_form.find("input", attrs={"name": "parent"})
        assert parent_field is not None and parent_field["value"] == parent
    for label in ("Add question", "Add sub-question"):
        assert f'class="button secondary">{label}</button>' in page
    assert not parent_card.select("select")

    response = client.post(
        f"/content/{night}/questions/{content.questions[1].ref}",
        data={
            "blob": _state(services, key, night).blob,
            "question": "How long?",
            "answer": "Two weeks.",
            "parent": "q1",
        },
        headers={"Accept": "application/json"},
    )
    assert response.status_code == 200
    updated = services.contents.result(key, _state(services, key, night))
    assert updated.questions[1].parent_ref == "q1"
    assert updated.questions[1].answer.text == "Two weeks."
    blob = _state(services, key, night).blob
    response = client.post(
        f"/content/{night}/questions/new",
        data={
            "blob": blob,
            "question": "More details?",
            "answer": "More.",
            "parent": updated.questions[1].ref,
        },
        headers={"Accept": "application/json"},
    )
    assert response.status_code == 400
    assert response.json["saved"] is False
    assert "one level" in response.json["message"]
    assert _state(services, key, night).blob == blob
    assert services.contents.result(key, _state(services, key, night)) == updated


@pytest.mark.parametrize("editable", [True, False])
def test_subquestions_render_once_inside_their_own_parent(
    generated, services, monkeypatch, editable
):
    client, key, night, _cash = generated

    def add(text, parent):
        response = client.post(
            f"/content/{night}/questions/new",
            data={
                "blob": _state(services, key, night).blob,
                "question": text,
                "answer": f"Answer to {text}",
                "parent": parent,
            },
        )
        assert response.status_code == 302
        return services.contents.result(key, _state(services, key, night)).questions[-1].ref

    other_parent = add("Where is the shelter?", "")
    children = {
        "q1": add("How long can I stay?", "q1"),
        other_parent: add("What are the opening hours?", other_parent),
    }
    if not editable:
        monkeypatch.setattr(
            "build_a_hia.workspace.drafting._content_status", lambda *_args: "outdated"
        )
    document = BeautifulSoup(client.get(f"/content/{night}").data, "html.parser")
    assert len(document.select("#questions > article.question")) == 2
    assert len(document.select("article.sub-question")) == 2
    assert not document.select("form form")
    for parent_ref, child_ref in children.items():
        parent = document.find("article", id=f"question-{parent_ref}")
        child = document.find("article", id=f"question-{child_ref}")
        assert parent is not None and child is not None
        assert child.find_parent("article") is parent
        assert len(parent.find_all("article", recursive=False)) == 1
        assert len(document.find_all(id=f"question-{child_ref}")) == 1
        assert (child.find("form", attrs={"data-content-autosave": True}) is not None) is editable


def test_gap_suggested_actions_are_read_only_on_overview_and_content(generated, services):
    client, key, night, _cash = generated
    gap = next(
        g
        for g in services.exports.plan(services.sessions.get_active(key)).gaps
        if g.subcategory_key == night
    )
    services.gaps.update(
        key, gap.id, status="resolved", action="Ask the desk", contact="Program coordinator"
    )

    for path in ("/content", f"/content/{night}"):
        page = client.get(path).data.decode()
        card = page.split(f'id="gap-{gap.id}"', 1)[1].split("</li>", 1)[0]
        assert "Ask the desk\nContact: Program coordinator" in card
        assert "Suggested contact" not in card
        assert 'name="suggested_contact"' not in card
        assert 'name="status"' not in card
        assert "badge" not in card
        assert "<details" not in card
        assert "<form" not in card
        assert "<textarea" not in card
        assert "<input" not in card
        assert "data-gap-autosave" not in card
        assert "data-save-status" not in card
        assert '<p class="gap-action">' in card
        assert card.index(gap.issue) < card.index("Suggested action")
        if "Sources:" in card:
            assert card.index("Suggested action") < card.index("Sources:")


@pytest.mark.parametrize("accept", ["text/html", "application/json"])
def test_suggested_action_updates_are_not_available(generated, services, accept):
    client, key, night, _cash = generated
    gap = next(
        gap
        for gap in services.exports.plan(services.sessions.get_active(key)).gaps
        if gap.subcategory_key == night
    )
    services.gaps.update(key, gap.id, status="open", action="Previous action", contact="")
    before = services.gaps.overrides(key)

    response = client.post(
        f"/gaps/{gap.id}",
        data={"suggested_action": "Changed action", "back": night},
        headers={"Accept": accept},
    )

    assert response.status_code == 404
    assert services.gaps.overrides(key) == before


def test_content_page_lists_all_gaps_in_one_section(generated, services):
    client, key, _night, _cash = generated

    page = client.get("/content").data.decode()

    assert "Review gaps" in page
    assert "Not supported by sources" in page
    assert "No supporting content" in page
    gaps = page[page.index('id="gaps"') :]
    assert "Night shelter ›" in gaps
    assert "Cash assistance ›" in gaps
    assert 'name="suggested_action"' not in gaps
    assert gaps.rindex('id="gap-') < gaps.index("Continue to download")
    assert "Continue to download" not in client.get("/content/status").data.decode()


def test_former_review_step_redirects_to_content(generated):
    client, _key, _night, _cash = generated

    response = client.get("/review")

    assert response.status_code == 302
    assert response.headers["Location"] == "/content"
    steps = BeautifulSoup(client.get("/content").data, "html.parser").select(".steps li")
    assert [step.get_text(strip=True) for step in steps] == [
        "Context",
        "Sources",
        "Structure",
        "Content",
        "Download",
    ]


def _badges(client) -> dict[str, list[str]]:
    page = BeautifulSoup(client.get("/content/status").data, "html.parser")
    badges = {}
    for row in page.select("li.source"):
        title = row.select_one(".source-title")
        assert title is not None
        badges[title.get_text(strip=True)] = [b.get_text(strip=True) for b in row.select(".badge")]
    return badges


def test_content_list_shows_one_status_from_generated_to_approved(generated, services):
    client, key, night, cash = generated
    assert _badges(client) == {
        "Night shelter": ["Needs review"],
        "Cash assistance": ["Needs review"],
    }

    content = services.contents.result(key, _state(services, key, night))
    client.post(
        f"/content/{night}/offers/{content.offers[0].id}",
        data=_offer_form(_state(services, key, night).blob),
    )
    assert _badges(client)["Night shelter"] == ["Needs review", "Edited"]

    _approve_all(client, services, key, night, cash)
    assert _badges(client) == {"Night shelter": ["Approved"], "Cash assistance": ["Left out"]}
    assert b"2 of 2 reviewed" in client.get("/content/status").data


def test_content_list_reloads_the_page_when_generation_finishes(generated):
    client, _key, _night, _cash = generated

    status = client.get("/content/status")

    assert b'data-poll-reload="true"' in status.data


def test_content_page_preserves_structure_order_after_approval(generated, services):
    client, key, night, _cash = generated
    page = client.get("/content").data
    assert page.index(b"Night shelter") < page.index(b"Cash assistance")
    assert b"0 of 2 reviewed" in page

    client.post(
        f"/content/{night}/approve",
        data={"blob": _state(services, key, night).blob, "disclosure_reviewed": "y"},
    )

    page = client.get("/content").data
    assert page.index(b"Night shelter") < page.index(b"Cash assistance")
    assert b"1 of 2 reviewed" in page
    assert "1 sub-category still needs review" in " ".join(page.decode().split())


@pytest.mark.parametrize("empty_first", [False, True])
@pytest.mark.parametrize("decision", ["keep", "drop"])
def test_approval_opens_the_next_review_then_scrolls_to_download(
    generated, services, empty_first, decision
):
    client, key, night, cash = generated
    submissions = [
        (f"/content/{night}/approve", {"blob": _state(services, key, night).blob}),
        (
            f"/content/{cash}/empty",
            {"blob": _state(services, key, cash).blob, "decision": decision},
        ),
    ]
    if empty_first:
        submissions.reverse()
    (first, first_data), (last, last_data) = submissions

    response = client.post(first, data=first_data)
    next_key = night if empty_first else cash
    assert response.headers["Location"] == f"/content/{next_key}"
    page = client.get(response.headers["Location"]).data
    assert b"This is the next sub-category to review." in page

    response = client.post(last, data=last_data)
    assert response.headers["Location"] == "/content#content-actions"
    page = BeautifulSoup(client.get("/content").data, "html.parser")
    assert page.find(id="content-actions") is not None
    subcategory_links = page.select("#content-status .source-title")
    assert [link.get_text(strip=True) for link in subcategory_links] == [
        "Night shelter",
        "Cash assistance",
    ]
    assert b"2 of 2 reviewed" in client.get("/content").data


def test_approval_returns_to_the_list_when_the_rest_is_not_reviewable(generated, services):
    client, key, night, cash = generated
    client.post(f"/content/{cash}/regenerate")

    response = client.post(
        f"/content/{night}/approve", data={"blob": _state(services, key, night).blob}
    )

    assert response.headers["Location"] == "/content#content-status"


def test_download_is_blocked_until_everything_is_reviewed(generated):
    client, _key, _night, _cash = generated

    page = client.get("/download")
    assert b"needs review and approval" in page.data
    assert b"decide whether to keep or drop" in page.data

    response = client.post("/download", follow_redirects=True)
    assert b"Latest snapshot" not in response.data


def test_downloads_are_separate_and_clean(generated, services):
    client, key, night, cash = generated
    _approve_all(client, services, key, night, cash)

    client.post("/download")
    latest = services.exports.latest(key)
    hia = client.get(f"/download/{latest['snapshot']}/hia.xlsx")
    internal = client.get(f"/download/{latest['snapshot']}/review-internal.xlsx")

    assert hia.status_code == internal.status_code == 200
    assert hia.headers["Content-Disposition"] == 'attachment; filename="hia.xlsx"'
    assert hia.headers["Cache-Control"] == "no-store"
    workbook = load_workbook(io.BytesIO(hia.data))
    values = " ".join(
        str(c.value)
        for ws in workbook.worksheets
        for r in ws.iter_rows()
        for c in r
        if c.value is not None
    )
    assert "Night shelter" in values
    assert "Cash assistance" not in values
    assert "S1-p1-c1" not in values and "Not supported" not in values
    review = load_workbook(io.BytesIO(internal.data))
    assert review.sheetnames == ["Gaps", "Evidence", "Sources"]
    assert latest["snapshot"] in review["Gaps"]["A1"].value


def test_conversion_warning_reaches_review_and_internal_workbook(generated, services):
    client, key, night, cash = generated
    [source] = services.sources.list_for(key)
    limitation = "Very little text could be read from page(s) 2."

    def warn(source):
        source.limitations = [limitation]
        return True

    services.sources.repository.mutate(key, source.id, warn)
    page = client.get("/content").data
    assert b"Conversion problem" in page and limitation.encode() in page
    plan = services.exports.plan(services.sessions.get_active(key))
    gap = next(gap for gap in plan.gaps if gap.kind == "conversion")
    assert gap.source_reference == source.reference
    _approve_all(client, services, key, night, cash)
    client.post("/download")
    latest = services.exports.latest(key)
    response = client.get(f"/download/{latest['snapshot']}/review-internal.xlsx")
    workbook = load_workbook(io.BytesIO(response.data))
    values = [str(cell.value) for row in workbook["Gaps"] for cell in row]
    assert any(limitation in value for value in values)
    assert "Conversion problem" in values


def test_model_scope_omissions_stay_out_of_public_export_but_keep_cited_gaps(
    generated, services, run_jobs
):
    client, key, night, cash = generated
    draft = {
        "offers": [offer(cited("Night shelter", "S1-p1-c1"))],
        "questions": [],
        "issues": [
            {
                "kind": "wrong_programme",
                "item": "Night shelter",
                "field": "need_to_know",
                "description": "250 TL monthly per person belongs to T-SUY Projesi.",
                "chunk_ids": ["S1-p1-c1"],
            }
        ],
    }
    client.post(f"/content/{night}/regenerate")
    run_jobs(draft)

    assert b"Information for another programme" in client.get("/content").data
    _approve_all(client, services, key, night, cash)
    assert _state(services, key, night).approved
    client.post("/download")
    latest = services.exports.latest(key)

    def values(filename):
        response = client.get(f"/download/{latest['snapshot']}/{filename}")
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data))
        return "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet
            for cell in row
            if cell.value is not None
        )

    public = values("hia.xlsx")
    assert "Night shelter" in public
    assert "250" not in public
    assert "T-SUY" not in public
    assert "ما قيمة المساعدة" not in public
    internal = values("review-internal.xlsx")
    assert "Information for another programme" in internal
    assert "S1-p1-c1" in internal
    assert "T-SUY Projesi" in internal
    assert "Find evidence for the intended programme" in internal


def test_model_contact_choices_and_conflict_gaps_reach_approved_exports(
    ready_client, services, run_jobs
):
    client = ready_client
    key = session_key_of(client)
    for phone, email, document_date, source_format in [
        ("213913935", "tracing@example.org", "2013-09-27", "pdf"),
        ("213 913 925", "rlf@example.org", "2025-10-23", "html"),
    ]:
        source = add_converted_source(
            services,
            key,
            [f"Family tracing: {phone}, {email}."],
            "Family tracing",
        )

        def metadata(source, document_date=document_date, source_format=source_format):
            source.document_date = document_date
            source.document_date_origin = "user"
            source.format = source_format
            return True

        services.sources.repository.mutate(key, source.id, metadata)
    services.sessions.approve_sources(key)
    proposal = {
        "categories": [
            {
                "name": "Family tracing",
                "description": "",
                "subcategories": [
                    {"name": name, "description": "", "chunk_ids": ["S2-p1-c1"]}
                    for name in ("Pesquisa de familiares", "Contacto em Portugal")
                ],
            }
        ],
        "notes": "",
    }
    client.post("/structure/propose", data={"instructions": ""})
    run_jobs(proposal)
    client.post("/structure/approve", data={"version": 1})
    client.post("/content/generate")
    scripts = []
    for name in ("Pesquisa de familiares", "Contacto em Portugal"):
        draft = {
            "offers": [
                offer(
                    cited(name, "S2-p1-c1"),
                    phone_numbers=[cited("213 913 925", "S3-p1-c1")],
                    emails=[cited("rlf@example.org", "S3-p1-c1")],
                )
            ],
            "questions": [
                question(
                    "q1",
                    f"How can I contact {name}?",
                    cited("Call 213 913 925.", "S3-p1-c1"),
                )
            ],
            "issues": [
                {
                    "kind": "conflict",
                    "item": name,
                    "field": field,
                    "description": f"The flyer lists {old}; the newer web page lists {new}.",
                    "chunk_ids": ["S2-p1-c1", "S3-p1-c1"],
                }
                for field, old, new in [
                    ("phone_numbers", "213913935", "213 913 925"),
                    ("emails", "tracing@example.org", "rlf@example.org"),
                ]
            ],
        }
        scripts.append(draft)
    run_jobs(*scripts)
    structure = services.structures.approved(key)[1]
    for _category, sub in structure.subcategories():
        state = _state(services, key, sub.key)
        assert state.status == "done"
        client.post(
            f"/content/{sub.key}/approve",
            data={"blob": state.blob, "disclosure_reviewed": "y"},
        )
        assert _state(services, key, sub.key).approved
    client.post("/download")
    latest = services.exports.latest(key)

    def values(filename):
        response = client.get(f"/download/{latest['snapshot']}/{filename}")
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data))
        return [
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet
            for cell in row
            if cell.value is not None
        ]

    public = values("hia.xlsx")
    assert public.count("213 913 925") == 2
    assert public.count("rlf@example.org") == 2
    assert public.count("Call 213 913 925.") == 2
    assert "213913935" not in "\n".join(public)
    assert "tracing@example.org" not in "\n".join(public)
    internal = "\n".join(values("review-internal.xlsx"))
    assert "213913935" in internal and "tracing@example.org" in internal
    assert "Conflicting sources" in internal
    assert "S2-p1-c1" in internal and "S3-p1-c1" in internal


def test_shared_contacts_in_single_call_drafts_reach_review_and_exports(
    ready_client, services, run_jobs
):
    client = ready_client
    key = session_key_of(client)
    add_converted_source(
        services,
        key,
        ["RLF offers Family messages and Trace the Face. Both services use the central RLF team."],
        "RLF services",
    )
    add_converted_source(
        services,
        key,
        [
            "RLF central team handles all RLF services, including family messages "
            "and Trace the Face. "
            "Call 213 913 925 or pesquisa@example.org. The central team is at Central Street 1."
        ],
        "Shared RLF contact",
    )
    services.sessions.approve_sources(key)
    client.post("/structure/propose", data={"instructions": ""})
    run_jobs(PROPOSAL)
    client.post("/structure/approve", data={"version": 1})
    client.post("/content/generate")
    scripts = []
    for name in ("Family messages", "Trace the Face"):
        draft = {
            "offers": [
                offer(
                    cited(name, "S2-p1-c1"),
                    phone_numbers=[cited("213 913 925", "S2-p1-c1", "S3-p1-c1")],
                    emails=[cited("pesquisa@example.org", "S2-p1-c1", "S3-p1-c1")],
                    address=cited("Central Street 1", "S2-p1-c1", "S3-p1-c1"),
                )
            ],
            "questions": [],
            "issues": [
                {
                    "kind": "missing",
                    "item": name,
                    "field": field,
                    "description": "The linked passages omit this contact.",
                    "chunk_ids": [],
                }
                for field in ("open_weekdays",)
            ],
        }
        scripts.append(draft)
    run_jobs(*scripts)
    structure = services.structures.approved(key)[1]
    for _category, sub in structure.subcategories():
        state = _state(services, key, sub.key)
        assert state.status == "done"
        content = services.contents.result(key, state)
        item = content.offers[0]
        assert item.emails[0].inherited_from == ""
        assert item.phone_numbers[0].inherited_from == ""
        assert {e.chunk_id for e in item.phone_numbers[0].evidence} == {
            "S2-p1-c1",
            "S3-p1-c1",
        }
        assert item.address.inherited_from == ""
        assert item.open_weekdays.text == ""
        assert [issue.field for issue in content.issues] == ["open_weekdays"]
        page = client.get(f"/content/{sub.key}").data
        assert b"Inherited programme contact" not in page
        assert b"213 913 925" in page
        client.post(
            f"/content/{sub.key}/approve",
            data={"blob": state.blob, "disclosure_reviewed": "y"},
        )
        assert _state(services, key, sub.key).approved
    client.post("/download")
    latest = services.exports.latest(key)

    def values(filename):
        response = client.get(f"/download/{latest['snapshot']}/{filename}")
        assert response.status_code == 200
        workbook = load_workbook(io.BytesIO(response.data))
        return "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet
            for cell in row
            if cell.value is not None
        )

    public = values("hia.xlsx")
    assert public.count("pesquisa@example.org") == 2
    assert public.count("213 913 925") == 2
    assert "Inherited programme contact" not in public
    internal = values("review-internal.xlsx")
    assert "Inherited programme contact" not in internal
    assert "S2-p1-c1" in internal
    assert "S3-p1-c1" in internal


def test_gap_noise_is_reduced_from_26_issues_to_four_gaps_in_review_and_export(
    ready_client, services, run_jobs
):
    client = ready_client
    key = session_key_of(client)
    recent = services.sources.repository.list_for(key)[0]

    def set_date(source):
        source.document_date = (date.today() - timedelta(days=364)).isoformat()
        source.document_date_origin = "user"
        return True

    services.sources.repository.mutate(key, recent.id, set_date)
    undated = add_converted_source(services, key, ["Service contact."], "Undated page")
    old = add_converted_source(services, key, ["Old service guidance."], "Old guide")

    def set_old_date(source):
        source.document_date = (date.today() - timedelta(days=366)).isoformat()
        source.document_date_origin = "user"
        return True

    services.sources.repository.mutate(key, old.id, set_old_date)
    services.sessions.approve_sources(key)
    client.post("/structure/propose", data={"instructions": ""})
    run_jobs(PROPOSAL)
    client.post("/structure/approve", data={"version": 1})
    client.post("/content/generate")
    drafts = []
    for name in ("Night shelter", "Cash assistance"):
        issues = (
            [
                {
                    "kind": "missing",
                    "item": name,
                    "field": field,
                    "description": "Optional field not given.",
                    "chunk_ids": [],
                }
                for field in ("address", "open_weekdays", "open_weekend")
            ]
            + [
                {
                    "kind": "outdated",
                    "item": name,
                    "field": field,
                    "description": "Possibly outdated.",
                    "chunk_ids": [chunk],
                }
                for field in ("description", "phone_numbers", "address")
                for chunk in ("S1-p1-c1", "S2-p1-c1", "S3-p1-c1")
            ]
            + [
                {
                    "kind": "conflict",
                    "item": name,
                    "field": "need_to_know",
                    "description": "Eligibility differs between sources.",
                    "chunk_ids": ["S1-p1-c1"],
                }
            ]
        )
        drafts.append(
            {"offers": [offer(cited(name, "S1-p1-c1"))], "questions": [], "issues": issues}
        )
    run_jobs(*drafts)
    workspace = services.sessions.get_active(key)
    plan = services.exports.plan(workspace)
    assert sum(len(sub.content.issues) for sub in plan.subcategories) == 26
    assert len(plan.gaps) == 4
    assert sorted(gap.kind for gap in plan.gaps) == ["conflict", "conflict", "outdated", "undated"]
    overview = client.get("/content").data.decode()
    assert "Conflicting sources: 2" in overview
    assert "Possibly outdated: 1" in overview
    assert "Source date unknown: 1" in overview
    assert "Optional field not given" not in overview
    assert f"src-{recent.id}-date" not in overview
    date_gap = f"src-{undated.id}-date"
    expected_action = next(gap.suggested_action for gap in plan.gaps if gap.id == date_gap)
    for sub in plan.subcategories:
        page = client.get(f"/content/{sub.key}").data.decode()
        assert "Eligibility differs" in page
        assert "Possibly outdated." not in page
        assert "Optional field not given" not in page
        client.post(
            f"/content/{sub.key}/approve",
            data={
                "blob": sub.state.blob,
                "disclosure_reviewed": "y",
            },
        )
    client.post("/download")
    snapshot = services.exports.latest(key)["snapshot"]
    response = client.get(f"/download/{snapshot}/review-internal.xlsx")
    workbook = load_workbook(io.BytesIO(response.data))
    rows = list(workbook["Gaps"].iter_rows(min_row=3, values_only=True))
    assert len(rows) == 4
    unknown_date = next(row for row in rows if row[5] == "Source date unknown")
    assert unknown_date[-3:] == (expected_action, None, "Open")


def test_downloads_are_isolated_and_replaced(generated, services, app):
    client, key, night, cash = generated
    _approve_all(client, services, key, night, cash)
    client.post("/download")
    first = services.exports.latest(key)["snapshot"]

    other = app.test_client()
    other.post(
        "/context",
        data={
            "country": "X",
            "situation": "Y",
            "target_group": "Z",
            "locations": "W",
            "output_language": "en",
        },
    )
    assert other.get(f"/download/{first}/hia.xlsx").status_code == 302

    client.post("/download")
    assert client.get(f"/download/{first}/hia.xlsx").status_code == 302
    assert client.get("/download/zzz/hia.xlsx").status_code == 404
    assert client.get(f"/download/{first}/secrets.txt").status_code == 404


def test_regeneration_resets_review(generated, services, run_jobs):
    client, key, night, cash = generated
    _approve_all(client, services, key, night, cash)

    client.post(f"/content/{night}/regenerate")
    run_jobs({"offers": [offer(cited("Night shelter", "S1-p1-c1"))], "questions": [], "issues": []})

    assert not _state(services, key, night).approved
    assert b"needs review and approval" in client.get("/download").data
