import io
import re
import zipfile
from dataclasses import replace
from datetime import UTC, datetime

import pytest
from factories import cited, question
from openpyxl import load_workbook

from build_a_hia.services import workbook as workbook_module
from build_a_hia.services.assembly import assemble
from build_a_hia.services.checks import ChunkIndex, ChunkRef, check_draft
from build_a_hia.services.content import (
    Evidence,
    FieldValue,
    GeneratedContent,
    Issue,
    Offer,
    Question,
    SubcategoryDraft,
)
from build_a_hia.services.context import ProjectContext
from build_a_hia.services.export import _formula_errors, _structural_errors
from build_a_hia.services.review import collect_gaps
from build_a_hia.services.sources import Source, SourceKind, SourceStatus
from build_a_hia.services.structure import Node, Structure
from build_a_hia.services.workbook import (
    SHEETS,
    TemplateError,
    build_hia_workbook,
    build_review_workbook,
    load_template,
)

EXPORTED_AT = datetime(2026, 10, 7, 12, 30, tzinfo=UTC)
CONTEXT = ProjectContext(
    country="Ukraine",
    situation="x",
    target_group="y",
    locations="z",
    output_language="ar",
    locale_dir="rtl",
)
EVIDENCE = Evidence(
    chunk_id="S1-p1-c1",
    source_id="x",
    source_ref="S1",
    source_title="Internal guide",
    page=1,
    passage="Secret internal passage",
)


def _categories():
    content = GeneratedContent(
        offers=[
            Offer(
                id="o1",
                name=FieldValue(
                    text='=HYPERLINK("http://evil.example","click")', evidence=[EVIDENCE]
                ),
                description=FieldValue(
                    text="Free beds.", evidence=[EVIDENCE], edited=True, original="Beds."
                ),
                phone_numbers=[FieldValue(text="0800 1234"), FieldValue(text="+380 44 123 4567")],
                open_weekdays=FieldValue(text="Mon-Fri 9:00-17:00"),
            )
        ],
        questions=[
            Question(ref="q1", question="Who can stay?", answer=FieldValue(text="Anyone.")),
            Question(
                ref="q2", parent_ref="q1", question="How long?", answer=FieldValue(text="@SUM(1)")
            ),
        ],
        issues=[
            Issue(
                id="i1",
                kind="missing",
                item_ref="o1",
                field="description",
                description="No description.",
            )
        ],
        generated_on="2026-10-07",
        prompt_version="t",
    )
    structure = Structure(
        categories=[
            Node(
                key="c1",
                name="Shelter",
                description="Places",
                children=[Node(key="s1", name="Beds")],
            )
        ]
    )
    return assemble(structure, {"s1": content})


def _all_values(data: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(data))
    return [
        str(cell.value)
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]


@pytest.fixture
def hia_bytes() -> bytes:
    return build_hia_workbook(_categories(), CONTEXT, EXPORTED_AT)


def test_text_that_would_run_as_a_formula_when_pasted_blocks_export():
    errors = _formula_errors(_categories())

    assert len(errors) == 1
    assert errors[0].startswith("Offer 2 (Name) starts like a formula")


def test_hia_workbook_keeps_the_template_tabs_and_headers(hia_bytes):
    exported = load_workbook(io.BytesIO(hia_bytes))
    template = load_template()

    assert tuple(exported.sheetnames) == SHEETS
    assert all(ws.sheet_state == "visible" for ws in exported.worksheets)
    for sheet in SHEETS:
        assert [c.value for c in exported[sheet][1]] == [c.value for c in template[sheet][1]]
    assert dict(exported.defined_names).keys() == dict(template.defined_names).keys()


def test_hia_rows_are_literal_values_directly_after_the_headers(hia_bytes):
    exported = load_workbook(io.BytesIO(hia_bytes))
    categories = exported["Categories"]
    offers = exported["Offers"]
    questions = exported["Q&As"]

    assert [c.value for c in categories[2][:5]] == [2, "Show", "shelter", "Shelter", "Places"]
    assert categories["A3"].value is None
    assert [c.value for c in exported["Sub-Categories"][2][:6]] == [
        "Shelter",
        2,
        2,
        "Show",
        "beds",
        "Beds",
    ]
    assert [offers[f"{c}2"].value for c in "ABCDE"] == ["Beds", 2, 2, 2, "Show"]
    assert offers["J2"].value == "0800 1234\n+380 44 123 4567"
    assert offers["N2"].value == "Mon-Fri 9:00-17:00"
    assert [questions[f"{c}2"].value for c in "FGHIJK"] == [
        "who-can-stay",
        None,
        "Who can stay?",
        "Anyone.",
        "2026-10-07",
        "No",
    ]
    assert questions["G3"].value == "who-can-stay"


def test_generated_chains_export_only_top_level_parents_in_the_same_subcategory():
    structure = Structure(
        categories=[
            Node(
                key="c1",
                name="Help",
                children=[Node(key="s1", name="First"), Node(key="s2", name="Second")],
            )
        ]
    )
    index = ChunkIndex([ChunkRef("S1-p1-c1", "s1", "S1", "Guide", 1, "", "Supported answer.")])
    answer = cited("Supported answer.", "S1-p1-c1")
    contents = {}
    for key, questions in {
        "s1": [
            question("q3", "Grandchild?", answer, "q2"),
            question("q2", "Child?", answer, "q1"),
            question("q1", "Root?", answer),
        ],
        "s2": [question("q4", "Other?", answer, "q1")],
    }.items():
        contents[key] = check_draft(
            SubcategoryDraft.model_validate({"offers": [], "questions": questions, "issues": []}),
            index,
            generated_on="2026-10-07",
            prompt_version="t",
        )
    categories = assemble(structure, contents)
    assert _structural_errors(categories) == []
    exported = load_workbook(io.BytesIO(build_hia_workbook(categories, CONTEXT, EXPORTED_AT)))
    rows = [
        row for row in exported["Q&As"].iter_rows(min_row=2, values_only=True) if row[5] is not None
    ]
    assert len(rows) == 4
    by_slug = {row[5]: row for row in rows}
    for row in rows:
        if row[6] is not None:
            parent = by_slug[row[6]]
            assert parent[1] == row[1]
            assert parent[6] is None
    assert {row[7]: row[6] for row in rows} == {
        "Root?": None,
        "Child?": "root",
        "Grandchild?": "root",
        "Other?": None,
    }


@pytest.mark.parametrize("parent_kind", ["nested", "self", "cross_subcategory", "missing"])
def test_export_validation_rejects_invalid_question_parents(parent_kind):
    categories = _categories()
    root, child = categories[0].subcategories[0].questions
    if parent_kind == "nested":
        categories[0].subcategories[0].questions.append(
            replace(child, slug="grandchild", parent_slug=child.slug, depth=2)
        )
    elif parent_kind == "self":
        child.parent_slug = child.slug
    elif parent_kind == "cross_subcategory":
        categories[0].subcategories[0].questions.remove(root)
        other = assemble(
            Structure(
                categories=[Node(key="c2", name="Other", children=[Node(key="s2", name="Other")])]
            ),
            {},
        )[0].subcategories[0]
        other.id = 3
        other.questions.append(root)
        categories[0].subcategories.append(other)
    else:
        child.parent_slug = "nonexistent"

    errors = _structural_errors(categories)

    assert any(
        ("more than one level" if parent_kind in {"nested", "self"} else "unknown parent") in error
        for error in errors
    )


@pytest.mark.parametrize("sheet", workbook_module.CONTRACT)
def test_generated_content_is_visible_by_default(hia_bytes, sheet):
    exported = load_workbook(io.BytesIO(hia_bytes))
    contract = workbook_module.CONTRACT[sheet]
    column = next(column for column, marker in contract.headers.items() if marker == "#VISIBLE")
    name_column = (
        "H"
        if sheet == "Q&As"
        else next(column for column, marker in contract.headers.items() if marker == "#NAME")
    )
    rows = [
        row
        for row in range(2, exported[sheet].max_row + 1)
        if exported[sheet][f"{name_column}{row}"].value is not None
    ]

    assert rows
    assert all(exported[sheet][f"{column}{row}"].value == "Show" for row in rows)


@pytest.mark.parametrize("sheet", workbook_module.CONTRACT)
def test_hidden_demo_rows_are_not_exported(hia_bytes, sheet):
    exported = load_workbook(io.BytesIO(hia_bytes))

    assert not any(
        cell.value in ("Hidden", "hidden", "example-question")
        for row in exported[sheet].iter_rows(min_row=2)
        for cell in row
    )


def test_empty_workbook_has_no_demo_rows():
    exported = load_workbook(io.BytesIO(build_hia_workbook([], CONTEXT, EXPORTED_AT)))

    for sheet in workbook_module.CONTRACT:
        assert all(
            cell.value is None for row in exported[sheet].iter_rows(min_row=2) for cell in row
        )


def test_text_never_becomes_a_formula(hia_bytes):
    exported = load_workbook(io.BytesIO(hia_bytes))

    name = exported["Offers"]["G2"]
    answer = exported["Q&As"]["I3"]
    assert name.value.startswith("=HYPERLINK") and name.data_type == "s"
    assert answer.value == "@SUM(1)" and answer.data_type == "s"
    with zipfile.ZipFile(io.BytesIO(hia_bytes)) as archive:
        sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet")]
        xml = "".join(archive.read(n).decode("utf-8") for n in sheets)
    formulas = re.findall(r"<f[^>]*>(.*?)</f>", xml)
    assert not any("HYPERLINK" in formula or "SUM" in formula for formula in formulas)


def test_demo_content_and_internal_data_are_not_exported(hia_bytes):
    values = "\n".join(_all_values(hia_bytes))

    # The Referral Page "#EXAMPLE value" column is template guidance and is kept;
    # its demo values (column B) are covered by the referral test below.
    for leaked in (
        "example.com",
        "Test/Example",
        "Category Alpha",
        "Secret internal passage",
        "S1-p1-c1",
        "Internal guide",
        "No address.",
        "Beds.",
        "GOOGLETRANSLATE",
    ):
        assert leaked not in values


def test_referral_page_contains_only_reviewed_settings(hia_bytes):
    referral = load_workbook(io.BytesIO(hia_bytes))["Referral Page"]
    filled = {
        str(referral[f"A{row}"].value).rsplit("\n", 1)[-1]: referral[f"B{row}"].value
        for row in range(2, referral.max_row + 1)
        if referral[f"B{row}"].value is not None
    }

    assert filled == {
        "#timestamp.last-updated": "2026-10-07T12:30:00.000Z",
        "#locale.dir": "rtl",
        "#locale.language": "ar",
    }
    assert referral["C13"].value


def test_workbook_metadata_is_neutral(hia_bytes):
    exported = load_workbook(io.BytesIO(hia_bytes))

    assert exported.properties.creator == "Build a HIA"
    assert exported.properties.lastModifiedBy is None
    assert not any(c.comment for ws in exported.worksheets for row in ws.iter_rows() for c in row)


def test_capacity_is_enforced():
    structure = Structure(
        categories=[
            Node(key=f"c{n}", name=f"Category {n}", children=[Node(key=f"s{n}", name=f"Sub {n}")])
            for n in range(100)
        ]
    )

    with pytest.raises(TemplateError):
        build_hia_workbook(assemble(structure, {}), CONTEXT, EXPORTED_AT)


def test_modified_template_is_rejected(monkeypatch):
    monkeypatch.setattr(workbook_module, "TEMPLATE_SHA256", "0" * 64)

    with pytest.raises(TemplateError):
        load_template()


def test_review_workbook_has_internal_tabs_and_stable_gap_columns():
    categories = _categories()
    sources = [
        Source(
            id="src1",
            number=1,
            kind=SourceKind.FILE,
            title="Internal guide",
            status=SourceStatus.CONVERTED,
            created_at="",
            limitations=["Page 3 blank."],
        )
    ]
    gaps = collect_gaps(categories, sources, {})

    data = build_review_workbook(
        categories=categories,
        gaps=gaps,
        sources=sources,
        approved={"s1"},
        snapshot="abc123",
        exported_at=EXPORTED_AT,
    )

    review = load_workbook(io.BytesIO(data))
    assert review.sheetnames == ["Gaps", "Evidence", "Sources"]
    for ws in review.worksheets:
        assert "INTERNAL REVIEW ONLY" in ws["A1"].value and "abc123" in ws["A1"].value
    assert [c.value for c in review["Gaps"][2]] == [
        "Sub-Category ID",
        "Sub-Category",
        "Item Type",
        "Item ID/Slug",
        "Field",
        "Issue Type",
        "Issue",
        "Source Reference",
        "Suggested Action",
        "Suggested Contact",
        "Status",
    ]
    gap_rows = [[c.value for c in row] for row in review["Gaps"].iter_rows(min_row=3)]
    assert gap_rows[0][:7] == [
        2,
        "Beds",
        "Offer",
        "2 / hyperlink-http-evil-example-click",
        "description",
        "Missing information",
        "No description.",
    ]
    evidence = [[c.value for c in row] for row in review["Evidence"].iter_rows(min_row=3)]
    description = next(row for row in evidence if row[4] == "Description")
    assert description[8] == "Secret internal passage"
    assert description[9] == "Approved"
    assert description[10] == "Edited by reviewer. Original: Beds."
    assert review["Gaps"]["G3"].data_type == "s"
    assert [c.value for c in review["Sources"][3]][:3] == ["S1", "Internal guide", None]
