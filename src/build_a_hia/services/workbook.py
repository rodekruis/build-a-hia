"""Build the two downloads: the publishable HIA workbook and the internal review workbook.

The HIA workbook starts from a pinned copy of the inspected demo workbook. Demo content is
cleared, including the hidden scaffolding rows, and generated rows are written as literal
values from row 2. Every text cell is stored as a string so it can never become a formula.
"""

import hashlib
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.packaging.core import DocumentProperties
from openpyxl.packaging.custom import CustomPropertyList
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from .assembly import HIGHLIGHT, VISIBLE, AssembledCategory
from .content import FIELD_LABELS, OFFER_LIST_FIELDS, OFFER_TEXT_FIELDS, FieldValue
from .context import ProjectContext
from .review import Gap, evidence_reference
from .sources import Source

TEMPLATE_VERSION = "hia-template-2026-10-06"
TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "hia_template" / f"{TEMPLATE_VERSION}.xlsx"
TEMPLATE_SHA256 = "4a8fd92a1224ca15e914c6c06836e78af5c7edd3af6e666edcac81fdbf623959"
SHEETS = (
    "Referral Page",
    "Help",
    "Options",
    "Categories",
    "Sub-Categories",
    "Offers",
    "Q&As",
    "Chat",
)
DATA_START_ROW = 2
MAX_CELL_LENGTH = 32_767
CREATOR = "Build a HIA"


@dataclass(frozen=True)
class SheetContract:
    """What the pinned template guarantees for one content tab.

    Attributes:
        headers: Expected header marker (e.g. "#ID") per column letter in row 1.
        warning_column: 1-based index of the template's warning column.
        last_row: Last row covered by the template's named and validated ranges.
    """

    headers: dict[str, str]
    warning_column: int
    last_row: int

    @property
    def capacity(self) -> int:
        """Number of data rows the tab can hold."""
        return self.last_row - DATA_START_ROW + 1


# Column markers verified in the pinned template; last rows match its named/validated ranges.
CONTRACT = {
    "Categories": SheetContract(
        headers={
            "A": "#ID",
            "B": "#VISIBLE",
            "C": "#SLUG",
            "D": "#NAME",
            "E": "#DESCRIPTION",
            "F": "#ICON",
        },
        warning_column=7,
        last_row=100,
    ),
    "Sub-Categories": SheetContract(
        headers={
            "B": "#CATEGORY",
            "C": "#ID",
            "D": "#VISIBLE",
            "E": "#SLUG",
            "F": "#NAME",
            "G": "#DESCRIPTION",
            "H": "#ICON",
        },
        warning_column=9,
        last_row=99,
    ),
    "Offers": SheetContract(
        headers={
            "B": "#SUBCATEGORY",
            "C": "#CATEGORY",
            "D": "#ID",
            "E": "#VISIBLE",
            "F": "#SLUG",
            "G": "#NAME",
            "H": "#DESCRIPTION",
            "I": "#ICON",
            "J": "#PHONENUMBERS",
            "K": "#EMAILS",
            "L": "#WEBURLS",
            "M": "#ADDRESS",
            "N": "#OPENWEEK",
            "O": "#OPENWEEKEND",
            "P": "#NEEDTOKNOW",
            "Q": "#MOREINFO",
            "R": "#CHAPTER",
        },
        warning_column=19,
        last_row=99,
    ),
    "Q&As": SheetContract(
        headers={
            "B": "#SUBCATEGORY",
            "C": "#CATEGORY",
            "D": "#VISIBLE",
            "F": "#SLUG",
            "G": "#PARENT",
            "H": "#QUESTION",
            "I": "#ANSWER",
            "J": "#UPDATED",
            "K": "#HIGHLIGHT",
        },
        warning_column=13,
        last_row=99,
    ),
}
REFERRAL_KEYS = ("#locale.dir", "#locale.language", "#timestamp.last-updated")


class TemplateError(Exception):
    """Raised when the pinned HIA template or the data does not fit the template contract."""


def _text(value: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_LENGTH]


def put(ws: Worksheet, coordinate: str, value: str | int | None) -> None:
    """Write a literal value to a cell; text can never become a formula.

    Text is stored as a string cell with text format, without characters Excel rejects and
    cut to the cell limit. None and "" clear the value.

    Args:
        ws: Worksheet to write to.
        coordinate: Cell coordinate, e.g. "B3".
        value: Text, integer, or None.
    """
    cell = ws[coordinate]
    if value is None or value == "":
        cell.value = None
    elif isinstance(value, int):
        cell.value = value
    else:
        cell.value = _text(value)
        # Stored as a plain string even when it starts with "=", "+", "-" or "@".
        cell.data_type = "s"
        cell.number_format = "@"


def load_template() -> Workbook:
    """Load the pinned HIA template and verify it against the contract.

    Returns:
        The template workbook, still with its demo content.

    Raises:
        TemplateError: If the checksum, the tabs or a header marker do not match.
    """
    data = TEMPLATE_PATH.read_bytes()
    if hashlib.sha256(data).hexdigest() != TEMPLATE_SHA256:
        raise TemplateError("The pinned HIA template file does not match its checksum.")
    workbook = load_workbook(io.BytesIO(data))
    if tuple(workbook.sheetnames) != SHEETS:
        raise TemplateError("The pinned HIA template does not have the expected tabs.")
    for sheet, contract in CONTRACT.items():
        ws = workbook[sheet]
        for column, marker in contract.headers.items():
            header = str(ws[f"{column}1"].value or "")
            if not header.rstrip().endswith(marker):
                raise TemplateError(f"Unexpected header in {sheet}!{column}1.")
    return workbook


def _reset_content_sheet(ws: Worksheet, contract: SheetContract) -> None:
    for merged in [str(item) for item in cast(Any, ws.merged_cells).ranges]:
        ws.unmerge_cells(merged)
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, contract.last_row)):
        for cell in row:
            _clear(cell)


def _clear(cell: Any) -> None:
    # Hyperlinks and comments keep demo URLs and notes even when the value is cleared.
    cell.value = None
    cell.hyperlink = None
    cell.comment = None


def _reset_referral_page(ws: Worksheet, context: ProjectContext, exported_at: datetime) -> None:
    values = {
        "#locale.dir": context.locale_dir,
        "#locale.language": context.locale_language,
        "#timestamp.last-updated": exported_at.strftime("%Y-%m-%dT%H:%M:%S.000Z"),
    }
    found: set[str] = set()
    for row in range(2, ws.max_row + 1):
        for column in ("B", "E", "F"):
            _clear(ws[f"{column}{row}"])
        key = str(ws[f"A{row}"].value or "").rstrip().rsplit("\n", 1)[-1].strip()
        if key in values:
            put(ws, f"B{row}", values[key])
            found.add(key)
    missing = set(REFERRAL_KEYS) - found
    if missing:
        raise TemplateError(f"Referral Page is missing keys: {', '.join(sorted(missing))}.")


def _join(values: list[FieldValue]) -> str:
    return "\n".join(value.text for value in values)


def build_hia_workbook(
    categories: list[AssembledCategory], context: ProjectContext, exported_at: datetime
) -> bytes:
    """Build the publishable HIA workbook from the pinned template.

    Demo content and hidden scaffolding rows are cleared, Referral Page settings
    are filled in, and rows are written as literal values from row 2. It contains no
    evidence, sources or review notes.

    Args:
        categories: Assembled rows, without dropped sub-categories.
        context: Project context with the locale settings.
        exported_at: Export time (UTC), used for #timestamp.last-updated and the file
            properties.

    Returns:
        The `.xlsx` file contents.

    Raises:
        TemplateError: If the template does not match its contract or a tab has too many
            rows.
    """
    workbook = load_template()
    for sheet, contract in CONTRACT.items():
        _reset_content_sheet(workbook[sheet], contract)
    _reset_referral_page(workbook["Referral Page"], context, exported_at)

    rows = dict.fromkeys(CONTRACT, DATA_START_ROW)

    def next_row(sheet: str) -> int:
        row = rows[sheet]
        if row > CONTRACT[sheet].last_row:
            raise TemplateError(f"Too many rows for the {sheet} tab.")
        rows[sheet] = row + 1
        return row

    ws_categories = workbook["Categories"]
    ws_subcategories = workbook["Sub-Categories"]
    ws_offers = workbook["Offers"]
    ws_questions = workbook["Q&As"]
    for category in categories:
        row = next_row("Categories")
        for column, value in zip(
            "ABCDEF",
            (category.id, VISIBLE, category.slug, category.name, category.description, ""),
            strict=True,
        ):
            put(ws_categories, f"{column}{row}", value)
        for sub in category.subcategories:
            row = next_row("Sub-Categories")
            for column, value in zip(
                "ABCDEFGH",
                (
                    category.name,
                    category.id,
                    sub.id,
                    VISIBLE,
                    sub.slug,
                    sub.name,
                    sub.description,
                    "",
                ),
                strict=True,
            ):
                put(ws_subcategories, f"{column}{row}", value)
            for item in sub.offers:
                offer = item.offer
                row = next_row("Offers")
                values = (
                    sub.name,
                    sub.id,
                    category.id,
                    item.id,
                    VISIBLE,
                    item.slug,
                    offer.name.text,
                    offer.description.text,
                    "",
                    _join(offer.phone_numbers),
                    _join(offer.emails),
                    _join(offer.web_urls),
                    offer.address.text,
                    offer.open_weekdays.text,
                    offer.open_weekend.text,
                    offer.need_to_know.text,
                    offer.more_info.text,
                    offer.chapter.text,
                )
                for column, value in zip("ABCDEFGHIJKLMNOPQR", values, strict=True):
                    put(ws_offers, f"{column}{row}", value)
            for item in sub.questions:
                row = next_row("Q&As")
                values = (
                    sub.name,
                    sub.id,
                    category.id,
                    VISIBLE,
                    "",
                    item.slug,
                    item.parent_slug,
                    item.question.question,
                    item.question.answer.text,
                    item.updated,
                    HIGHLIGHT,
                )
                for column, value in zip("ABCDEFGHIJK", values, strict=True):
                    put(ws_questions, f"{column}{row}", value)

    _set_properties(workbook, exported_at, title=None)
    return _save(workbook)


def _set_properties(workbook: Workbook, exported_at: datetime, *, title: str | None) -> None:
    naive = exported_at.replace(tzinfo=None)
    workbook.properties = DocumentProperties(
        creator=CREATOR, title=title, created=naive, modified=naive, lastModifiedBy=None
    )
    workbook.custom_doc_props = CustomPropertyList()


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


INTERNAL_WARNING = (
    "INTERNAL REVIEW ONLY. Do not copy anything from this workbook into a live HIA sheet: it "
    "contains source references, passages and review notes. Snapshot {snapshot}, exported "
    "{exported} UTC."
)
GAP_COLUMNS = (
    "Sub-Category ID",
    "Sub-Category",
    "Item Type",
    "Item ID/Slug",
    "Field",
    "Issue Type",
    "Issue",
    "Source Reference",
    "Suggested Action",
    "Suggested Contact",
    "Status",
)
EVIDENCE_COLUMNS = (
    "Sub-Category ID",
    "Sub-Category",
    "Item Type",
    "Item ID/Slug",
    "Field",
    "Value",
    "Source Reference",
    "Source Title",
    "Supporting Passage",
    "Review Status",
    "Human Edit",
)
SOURCE_COLUMNS = (
    "Source",
    "Title",
    "Type",
    "File Name",
    "Address",
    "Final Address",
    "Document Date",
    "Date Origin",
    "Retrieved (UTC)",
    "Status",
    "Pages",
    "Passages",
    "Processing Notes",
)


def _table_sheet(
    workbook: Workbook, title: str, columns: tuple[str, ...], warning: str
) -> Worksheet:
    ws = workbook.create_sheet(title)
    put(ws, "A1", warning)
    ws["A1"].font = Font(bold=True, color="A63725")
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=index)
        put(ws, cell.coordinate, name)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = 24
    ws.freeze_panes = "A3"
    return ws


def _append(ws: Worksheet, values: tuple[str | int | None, ...]) -> None:
    row = ws.max_row + 1
    for index, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=index)
        put(ws, cell.coordinate, value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _evidence_rows(
    categories: list[AssembledCategory], approved: set[str]
) -> list[tuple[str | int | None, ...]]:
    rows: list[tuple[str | int | None, ...]] = []

    def add(
        sub_id: int,
        sub_name: str,
        item_type: str,
        item: str,
        field: str,
        value: FieldValue,
        status: str,
    ) -> None:
        human = f"Edited by reviewer. Original: {value.original}" if value.edited else ""
        if value.edited and not value.original:
            human = "Added by reviewer."
        if not value.evidence:
            rows.append(
                (sub_id, sub_name, item_type, item, field, value.text, "", "", "", status, human)
            )
        for evidence in value.evidence:
            rows.append(
                (
                    sub_id,
                    sub_name,
                    item_type,
                    item,
                    field,
                    value.text,
                    evidence_reference([evidence]),
                    evidence.source_title,
                    evidence.passage,
                    status,
                    human,
                )
            )

    for category in categories:
        for sub in category.subcategories:
            status = "Approved" if sub.key in approved else "Not approved"
            for item in sub.offers:
                label = f"{item.id} / {item.slug}"
                for field in OFFER_TEXT_FIELDS:
                    value = getattr(item.offer, field)
                    if value.text:
                        add(sub.id, sub.name, "Offer", label, FIELD_LABELS[field], value, status)
                for field in OFFER_LIST_FIELDS:
                    for value in getattr(item.offer, field):
                        add(sub.id, sub.name, "Offer", label, FIELD_LABELS[field], value, status)
            for item in sub.questions:
                question = FieldValue(
                    text=item.question.question, edited=item.question.question_edited
                )
                add(sub.id, sub.name, "Q&A", item.slug, "Question", question, status)
                add(sub.id, sub.name, "Q&A", item.slug, "Answer", item.question.answer, status)
    return rows


def build_review_workbook(
    *,
    categories: list[AssembledCategory],
    gaps: list[Gap],
    sources: list[Source],
    approved: set[str],
    snapshot: str,
    exported_at: datetime,
) -> bytes:
    """Build the internal review workbook with Gaps, Evidence and Sources tabs.

    Every tab starts with a warning that the workbook is not for publication.

    Args:
        categories: Assembled rows that are in the HIA workbook.
        gaps: Gaps with their reviewer status and notes.
        sources: Sources of the session.
        approved: Keys of the sub-categories whose content is approved.
        snapshot: Snapshot ID shown in the warning.
        exported_at: Export time (UTC).

    Returns:
        The `.xlsx` file contents.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    warning = INTERNAL_WARNING.format(
        snapshot=snapshot, exported=exported_at.strftime("%Y-%m-%d %H:%M")
    )

    gap_sheet = _table_sheet(workbook, "Gaps", GAP_COLUMNS, warning)
    for gap in gaps:
        _append(
            gap_sheet,
            (
                gap.subcategory_id,
                gap.subcategory,
                gap.item_type,
                gap.item,
                gap.field,
                gap.kind_label,
                gap.issue,
                gap.source_reference,
                gap.suggested_action,
                gap.suggested_contact,
                gap.status_label,
            ),
        )

    evidence_sheet = _table_sheet(workbook, "Evidence", EVIDENCE_COLUMNS, warning)
    for row in _evidence_rows(categories, approved):
        _append(evidence_sheet, row)

    source_sheet = _table_sheet(workbook, "Sources", SOURCE_COLUMNS, warning)
    for source in sources:
        notes = [source.error_message] if source.error_message else []
        notes += source.limitations
        _append(
            source_sheet,
            (
                source.reference,
                source.title,
                source.format,
                source.filename,
                source.url,
                source.final_url,
                source.document_date,
                source.document_date_origin,
                source.retrieved_at[:19].replace("T", " "),
                source.status.value,
                source.page_count or None,
                source.chunk_count or None,
                "\n".join(notes),
            ),
        )

    _set_properties(workbook, exported_at, title="Internal review - not for publication")
    return _save(workbook)
