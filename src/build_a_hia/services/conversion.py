"""Convert sources to markdown with Docling on CPU.

Follows the approach used in rodekruis/appeals-monitor: Docling is imported lazily, PDF
pages are counted with pypdfium2, PDFs are converted in small page ranges with page-based
timeouts, and full-page OCR retries pages with missing or sparse readable text.
"""

import io
import logging
import re
from collections import OrderedDict
from collections.abc import Callable
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path
from types import SimpleNamespace
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from bs4 import BeautifulSoup, Tag

from .chunking import ConvertedUnit, estimate_tokens
from .inspection import IMAGE_FORMATS, SourceError, SourceFormat
from .languages import language_label, resolve_language

logger = logging.getLogger(__name__)

TIMEOUT_PER_PAGE = 30.0
MIN_TIMEOUT = 60.0
MAX_TIMEOUT = 600.0
# Small page ranges avoid native out-of-memory errors (std::bad_alloc) on large PDFs.
PDF_CHUNK_PAGES = 5
MIN_READABLE_TOKENS = 30
# RapidOCR language tags supported by Docling's installed PP-OCR vocabularies.
OCR_LANGUAGES = {
    "af": "iso:af",
    "ar": "arabic",
    "az": "iso:az",
    "be": "iso:be",
    "bs": "iso:bs",
    "ca": "iso:ca",
    "cs": "iso:cs",
    "cy": "iso:cy",
    "da": "iso:da",
    "de": "iso:de",
    "el": "iso:el",
    "en": "iso:en",
    "es": "iso:es",
    "et": "iso:et",
    "eu": "iso:eu",
    "fi": "iso:fi",
    "fil": "iso:fil",
    "fr": "iso:fr",
    "ga": "iso:ga",
    "gl": "iso:gl",
    "hr": "iso:hr",
    "hu": "iso:hu",
    "id": "iso:id",
    "is": "iso:is",
    "it": "iso:it",
    "ja": "iso:ja",
    "kn": "iso:kn",
    "ko": "iso:ko",
    "ku": "iso:ku",
    "la": "iso:la",
    "lb": "iso:lb",
    "lt": "iso:lt",
    "lv": "iso:lv",
    "mi": "iso:mi",
    "ms": "iso:ms",
    "mt": "iso:mt",
    "nl": "iso:nl",
    "no": "iso:no",
    "oc": "iso:oc",
    "pl": "iso:pl",
    "pt": "iso:pt",
    "qu": "iso:qu",
    "rm": "iso:rm",
    "ro": "iso:ro",
    "ru": "iso:ru",
    "sk": "iso:sk",
    "sl": "iso:sl",
    "sq": "iso:sq",
    "sr-Latn": "iso:sr-Latn",
    "sv": "iso:sv",
    "sw": "iso:sw",
    "ta": "iso:ta",
    "te": "iso:te",
    "th": "iso:th",
    "tr": "iso:tr",
    "uk": "iso:uk",
    "uz": "iso:uz",
    "vi": "iso:vi",
    "zh": "iso:zh-Hans",
    "zh-Hant": "iso:zh-Hant",
}
OCR_PREFETCH_LANGUAGES = ("en", "ar", "be", "el", "kn", "ko", "ta", "te", "th")
_OCR_LANGUAGE_LABEL_OVERRIDES = {
    "ku": "Kurdish",
    "la": "Latin",
    "no": "Norwegian",
    "oc": "Occitan",
    "rm": "Romansh",
    "sr-Latn": "Serbian (Latin)",
    "zh": "Chinese (Simplified)",
    "zh-Hant": "Chinese (Traditional)",
}
MAX_CACHED_OCR_CONVERTERS = 4
_OCR_LANGUAGE_CODES_BY_LABEL = {
    label.casefold(): code
    for code in OCR_LANGUAGES
    for label in (language_label(code), _OCR_LANGUAGE_LABEL_OVERRIDES.get(code, ""))
    if label
}
_NON_TEXT = re.compile(r"<!--.*?-->|!\[[^\]]*\]\([^)]*\)", re.DOTALL)
_MARKDOWN_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_HEADING = re.compile(r"^ {0,3}#{1,6}\s+", re.MULTILINE)
_HIDDEN_STYLE = re.compile(
    r"^\s*(?:display\s*:\s*none|visibility\s*:\s*(?:hidden|collapse)|opacity\s*:\s*0)"
    r"\s*(?:!important\s*)?$",
    re.IGNORECASE,
)
_HTML_UI_MARKER = re.compile(
    r"(?:^|[-_\s])(?:flyout|modal|popup|pop-up|breadcrumbs?|newsletter|"
    r"cookie-(?:banner|consent|notice)|site-(?:menu|navigation))(?=$|[-_\s])",
    re.IGNORECASE,
)
_HTML_DECORATIVE_IMAGE = re.compile(
    r"(?:^|[-_.\s/])(?:logos?|icons?|ico|avatars?|spacers?)(?=$|[-_.\s/])", re.IGNORECASE
)
_HTML_IMAGE_TAG = re.compile(rb"<img[\s/>]", re.IGNORECASE)


def _html_image_src(tag: "Tag") -> str:
    for attribute in ("data-src", "data-lazy-src", "src"):
        value = tag.get(attribute)
        if (
            isinstance(value, str)
            and value.strip()
            and not value.strip().lower().startswith("data:")
        ):
            return value.strip()
    for attribute in ("data-srcset", "data-lazy-srcset", "srcset"):
        value = tag.get(attribute)
        if isinstance(value, str) and value.strip():
            candidate = value.split(",")[-1].strip().split()
            if candidate:
                return candidate[0]
    value = tag.get("src")
    return value.strip() if isinstance(value, str) else ""


def _html_is_decorative_image(tag: "Tag") -> bool:
    if _html_image_src(tag).lower().startswith("data:image/svg+xml"):
        return True
    if _html_roles(tag) & {"presentation", "none"}:
        return True
    if str(tag.get("aria-hidden", "")).strip().lower() == "true":
        return True
    classes = tag.get("class")
    class_names = classes if isinstance(classes, list) else [str(classes or "")]
    identity = " ".join(
        [str(tag.get("id", "")), *class_names, _html_image_src(tag).split("?", 1)[0]]
    )
    return bool(_HTML_DECORATIVE_IMAGE.search(identity))


def _html_roles(tag: "Tag") -> set[str]:
    role = tag.get("role")
    return set(role.lower().split()) if isinstance(role, str) else set()


def _html_is_main(tag: "Tag") -> bool:
    return tag.name == "main" or "main" in _html_roles(tag)


def _html_is_widget(tag: "Tag") -> bool:
    ui_roles = {"navigation", "banner", "contentinfo", "dialog", "alertdialog", "search"}
    if _html_roles(tag) & ui_roles:
        return True
    if str(tag.get("aria-modal", "")).strip().lower() == "true":
        return True
    classes = tag.get("class")
    class_names = classes if isinstance(classes, list) else [str(classes or "")]
    identity = " ".join([str(tag.get("id", "")), *class_names])
    return bool(_HTML_UI_MARKER.search(identity))


def _html_main_content(soup: "BeautifulSoup") -> None:
    """Keep all nonempty main landmarks, or articles when no main exists.

    Image-only landmarks count as nonempty: their images are not read, so such pages fail
    with a clear message instead of silently falling back to unrelated body text.
    Retain the body when no usable landmark exists (common for static FAQ fragments).
    Article headers/footers and static service sidebars inside main are content, not
    site chrome; do not discard them just because of their tag names.
    """
    candidates = [
        tag
        for tag in soup.find_all(True)
        if _html_is_main(tag) and (tag.get_text(strip=True) or tag.find("img") is not None)
    ]
    if not candidates:
        candidates = [
            tag
            for tag in soup.find_all("article")
            if tag.get_text(strip=True) or tag.find("img") is not None
        ]
    if not candidates:
        return
    selected = [
        tag for tag in candidates if not any(parent in candidates for parent in tag.parents)
    ]
    body = soup.body
    if body is None:
        body = soup.new_tag("body")
        (soup.html or soup).append(body)
    content = [tag.extract() for tag in selected]
    body.clear()
    for tag in content:
        body.append(tag)


def _normalize_html(data: bytes) -> bytes:
    """Retain static main content/FAQs while excluding site chrome and interactive UI."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    # Work inside out so removing a parent never leaves decomposed child tags to inspect.
    for tag in reversed(soup.find_all(True)):
        site_chrome = tag.name in {"header", "footer"} and not any(
            parent.name == "article" or _html_is_main(parent) for parent in tag.parents
        )
        if (
            tag.name in {"script", "noscript", "style", "template", "nav"}
            or (tag.name == "form" and tag.find(["input", "select", "textarea"]) is not None)
            or _html_is_widget(tag)
            or site_chrome
            or (tag.name == "img" and _html_is_decorative_image(tag))
        ):
            tag.decompose()
        elif tag.name in {"header", "footer"}:
            # Docling otherwise skips some text in headers/footers, even within articles.
            tag.name = "div"
    _html_main_content(soup)
    for tag in soup.find_all(True):
        tag.attrs.pop("hidden", None)
        tag.attrs.pop("aria-hidden", None)
        style = tag.get("style")
        if isinstance(style, str):
            tag["style"] = ";".join(
                declaration
                for declaration in style.split(";")
                if not _HIDDEN_STYLE.fullmatch(declaration)
            )
    # Docling accepts one wrapper around dt/dd, but not nested accordion wrappers.
    for description_list in soup.find_all("dl"):
        for wrapper in reversed(description_list.find_all("div")):
            if (
                wrapper.find_parent(["dl", "dt", "dd"]) is description_list
                and wrapper.find(["dt", "dd"]) is not None
            ):
                wrapper.unwrap()
        # Docling's nested-list path otherwise drops prose next to the inner list.
        if description_list.find("dl") is not None:
            entries = [
                tag
                for tag in description_list.find_all(["dt", "dd"])
                if tag.find_parent("dl") is description_list
            ]
            for entry in entries:
                entry.name = "h3" if entry.name == "dt" else "div"
            description_list.name = "div"
    return soup.encode("utf-8")


def _html_has_sparse_body(markdown: str) -> bool:
    lines = _NON_TEXT.sub("", markdown).splitlines()
    headings = [line for line in lines if _HEADING.match(line)]
    if len(headings) < 3:
        return False
    heading_tokens = _text_tokens("\n".join(headings))
    body_tokens = _text_tokens("\n".join(line for line in lines if not _HEADING.match(line)))
    return body_tokens < 30 and body_tokens * 4 < heading_tokens


def readable_text(markdown: str) -> str:
    """Return visible text for extraction checks, ignoring comments, images and markup."""
    text = _MARKDOWN_LINK.sub(r"\1", _NON_TEXT.sub("", markdown))
    return " ".join(re.sub(r"[#*`_~|>\-]+", " ", text).split())


def _text_tokens(markdown: str) -> int:
    text = readable_text(markdown)
    return estimate_tokens(text) if any(character.isalnum() for character in text) else 0


def ocr_language_label(language: str) -> str:
    """Return the user-facing name of a supported OCR language."""
    return _OCR_LANGUAGE_LABEL_OVERRIDES.get(language, language_label(language))


def validate_ocr_language(language: str) -> str:
    """Return a canonical supported OCR language from its name or code."""
    value = language.strip()
    code = next((code for code in OCR_LANGUAGES if code.casefold() == value.casefold()), None)
    if code is None:
        code = _OCR_LANGUAGE_CODES_BY_LABEL.get(value.casefold())
    if code is None:
        resolved = resolve_language(value)
        if resolved in OCR_LANGUAGES:
            code = resolved
    if code is None:
        raise SourceError(
            "invalid_ocr_language", "Choose a language from the supported OCR suggestions."
        )
    return code


@dataclass
class _PageConversion:
    texts: dict[int, str]
    partial: bool = False


ProgressCallback = Callable[[int, int], None]


class ConversionCancelledError(Exception):
    """Raised by a progress callback to stop a conversion that is no longer wanted."""


@dataclass
class ConversionOutput:
    """Converted text of a source.

    Attributes:
        units: Pages, sheets or the whole document, limited to units with readable text.
        limitations: User-facing notes on content that could not be converted.
    """

    units: list[ConvertedUnit]
    limitations: list[str] = field(default_factory=list)


def calculate_timeout(num_pages: int) -> float:
    """Return the conversion timeout in seconds for a number of pages, within fixed bounds."""
    return min(max(num_pages * TIMEOUT_PER_PAGE, MIN_TIMEOUT), MAX_TIMEOUT)


def page_chunks(num_pages: int, size: int = PDF_CHUNK_PAGES) -> list[tuple[int, int]]:
    """Split pages 1 to `num_pages` into inclusive ranges of at most `size` pages."""
    return [(start, min(start + size - 1, num_pages)) for start in range(1, num_pages + 1, size)]


def page_ranges(pages: list[int]) -> list[tuple[int, int]]:
    """Group page numbers into sorted, inclusive ranges of consecutive pages."""
    ranges: list[tuple[int, int]] = []
    for page in sorted(pages):
        if ranges and page == ranges[-1][1] + 1:
            ranges[-1] = (ranges[-1][0], page)
        else:
            ranges.append((page, page))
    return ranges


def describe_pages(pages: list[int]) -> str:
    """Format page numbers for users, e.g. ``1-3, 7``."""
    return ", ".join(
        str(start) if start == end else f"{start}-{end}" for start, end in page_ranges(pages)
    )


def count_pdf_pages(path: Path) -> int:
    """Return the number of pages of a PDF file."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(path)
    try:
        return len(document)
    finally:
        document.close()


@cache
def _docling() -> SimpleNamespace:
    from docling.datamodel.backend_options import HTMLBackendOptions
    from docling.datamodel.base_models import ConversionStatus, DocumentStream, InputFormat
    from docling.datamodel.pipeline_options import OcrMode, PdfPipelineOptions, RapidOcrOptions
    from docling.document_converter import (
        DocumentConverter,
        HTMLFormatOption,
        ImageFormatOption,
        PdfFormatOption,
    )

    return SimpleNamespace(
        ConversionStatus=ConversionStatus,
        DocumentConverter=DocumentConverter,
        DocumentStream=DocumentStream,
        HTMLBackendOptions=HTMLBackendOptions,
        HTMLFormatOption=HTMLFormatOption,
        ImageFormatOption=ImageFormatOption,
        InputFormat=InputFormat,
        RapidOcrOptions=RapidOcrOptions,
        OcrMode=OcrMode,
        PdfFormatOption=PdfFormatOption,
        PdfPipelineOptions=PdfPipelineOptions,
    )


class DoclingConverter:
    """Convert source files to markdown with Docling, reusing converters across calls."""

    def __init__(self) -> None:
        # Converters cache their loaded models, so reuse them across chunks.
        self._converters: OrderedDict[tuple[bool, float, bool, str], Any] = OrderedDict()

    def convert(
        self,
        path: Path,
        source_format: SourceFormat,
        progress: ProgressCallback,
        *,
        base_url: str = "",
        ocr_language: str = "en",
    ) -> ConversionOutput:
        """Convert a file into units of markdown.

        PDFs are converted in small page ranges, and pages with missing or sparse readable
        text are retried with full-page OCR. Other formats convert as a whole; images use OCR.

        Args:
            path: File to convert.
            source_format: Format detected by inspection.
            progress: Called with pages done and total pages; may raise to stop the conversion.
            base_url: Address of a fetched web page, used to make its relative links absolute.
                Nothing is fetched from it; text in its content images is not read.
            ocr_language: Language of scanned text, selected separately for each source.

        Returns:
            The units with readable text and any limitations.

        Raises:
            SourceError: ``no_text`` if no readable text was found, or ``conversion_failed``
                if a non-PDF file could not be converted; ``invalid_ocr_language`` for an
                unsupported recognition language.
        """
        ocr_language = validate_ocr_language(ocr_language)
        if source_format == SourceFormat.PDF:
            return self._convert_pdf(path, progress, ocr_language)
        output = self._convert_whole(path, source_format, progress, base_url, ocr_language)
        if source_format in IMAGE_FORMATS and any(
            _text_tokens(unit.markdown) < MIN_READABLE_TOKENS for unit in output.units
        ):
            output.limitations.append(
                f"Very little text could be read from this image "
                f"(under {MIN_READABLE_TOKENS} estimated tokens per page). "
                "Check the converted text and selected language."
            )
        return output

    def prefetch_models(self) -> None:
        """Download and load the PDF and image pipeline models before the first conversion."""
        docling = _docling()
        self._converter(paginated=True, timeout=MIN_TIMEOUT, with_ocr=False).initialize_pipeline(
            docling.InputFormat.PDF
        )
        # Each entry represents a distinct PP-OCR checkpoint family.
        for code in OCR_PREFETCH_LANGUAGES:
            converter = self._converter(
                paginated=True, timeout=MIN_TIMEOUT, with_ocr=True, ocr_language=code
            )
            converter.initialize_pipeline(docling.InputFormat.PDF)
            converter.initialize_pipeline(docling.InputFormat.IMAGE)
            self._converters.clear()

    def _converter(
        self, *, paginated: bool, timeout: float, with_ocr: bool, ocr_language: str = "en"
    ) -> Any:
        key = (paginated, timeout, with_ocr, OCR_LANGUAGES[ocr_language] if with_ocr else "")
        if key in self._converters:
            self._converters.move_to_end(key)
            return self._converters[key]
        docling = _docling()
        if paginated:
            options = docling.PdfPipelineOptions(document_timeout=timeout, do_ocr=with_ocr)
            if with_ocr:
                options.ocr_options = docling.RapidOcrOptions(
                    mode=docling.OcrMode.FULL_PAGE, lang=[OCR_LANGUAGES[ocr_language]]
                )
            converter = docling.DocumentConverter(
                allowed_formats=[docling.InputFormat.PDF, docling.InputFormat.IMAGE],
                format_options={
                    docling.InputFormat.PDF: docling.PdfFormatOption(pipeline_options=options),
                    docling.InputFormat.IMAGE: docling.ImageFormatOption(pipeline_options=options),
                },
            )
        else:
            converter = docling.DocumentConverter(
                allowed_formats=[docling.InputFormat.DOCX, docling.InputFormat.XLSX],
            )
        self._converters[key] = converter
        if with_ocr:
            ocr_keys = [cached_key for cached_key in self._converters if cached_key[2]]
            if len(ocr_keys) > MAX_CACHED_OCR_CONVERTERS:
                del self._converters[ocr_keys[0]]
        return converter

    @staticmethod
    def _html_converter(base_url: str) -> Any:
        docling = _docling()
        html_options = docling.HTMLBackendOptions(
            enable_remote_fetch=False,
            enable_local_fetch=False,
            fetch_images=False,
            render_page=False,
            # Site chrome is already removed; text before the first heading is content.
            infer_furniture=False,
            source_uri=base_url or None,
        )
        return docling.DocumentConverter(
            allowed_formats=[docling.InputFormat.HTML],
            format_options={
                docling.InputFormat.HTML: docling.HTMLFormatOption(backend_options=html_options),
            },
        )

    def _convert_pdf(
        self, path: Path, progress: ProgressCallback, ocr_language: str
    ) -> ConversionOutput:
        total = count_pdf_pages(path)
        progress(0, total)
        texts: dict[int, str] = {}
        retry: list[tuple[int, int]] = []
        partial: set[int] = set()
        done = 0

        for start, end in page_chunks(total):
            converted = self._convert_range(
                path, (start, end), with_ocr=False, ocr_language=ocr_language
            )
            if converted is None:
                retry.append((start, end))
            else:
                if converted.partial:
                    partial.update(range(start, end + 1))
                texts.update(
                    {page: text for page, text in converted.texts.items() if _text_tokens(text)}
                )
                sparse = [
                    page
                    for page in range(start, end + 1)
                    if _text_tokens(texts.get(page, "")) < MIN_READABLE_TOKENS
                ]
                retry.extend(page_ranges(sparse))
                done += end - start + 1 - len(sparse)
            progress(done, total)

        missing: list[int] = []
        ocr_failed: list[int] = []
        for start, end in retry:
            converted = self._convert_range(
                path, (start, end), with_ocr=True, ocr_language=ocr_language
            )
            if converted is None:
                ocr_failed.extend(range(start, end + 1))
            elif converted.partial:
                partial.update(range(start, end + 1))
            for page in range(start, end + 1):
                text = converted.texts.get(page, "") if converted is not None else ""
                if _text_tokens(text) > _text_tokens(texts.get(page, "")):
                    texts[page] = text
                if page not in texts:
                    missing.append(page)
            done += end - start + 1
            progress(done, total)

        if not texts:
            message = "No readable text was found in this document."
            if ocr_failed:
                message += " OCR could not complete. Try again or upload a text-based copy."
            else:
                message += " Check the selected language for scanned text or add a better copy."
            raise SourceError("no_text", message)
        limitations = []
        if missing:
            limitations.append(
                f"No text could be read from page(s) {describe_pages(missing)}. "
                "They may be blank or contain only images."
            )
        if partial:
            limitations.append(
                f"Conversion reported partial results for page(s) {describe_pages(list(partial))}. "
                "Check the converted text; some information may be unavailable."
            )
        sparse = [page for page, text in texts.items() if _text_tokens(text) < MIN_READABLE_TOKENS]
        if sparse:
            limitations.append(
                f"Very little text could be read from page(s) {describe_pages(sparse)} "
                f"(under {MIN_READABLE_TOKENS} estimated tokens per page). "
                "Check the converted text and selected language; "
                "some information may be unavailable."
            )
        if ocr_failed:
            limitations.append(
                f"OCR could not complete for page(s) {describe_pages(ocr_failed)}. "
                "Any readable native text was kept; scanned information may be unavailable."
            )
        units = [ConvertedUnit(page=page, markdown=texts[page]) for page in sorted(texts)]
        return ConversionOutput(units=units, limitations=limitations)

    def _convert_range(
        self, path: Path, page_range: tuple[int, int], *, with_ocr: bool, ocr_language: str = "en"
    ) -> _PageConversion | None:
        docling = _docling()
        timeout = calculate_timeout(page_range[1] - page_range[0] + 1)
        try:
            converter = self._converter(
                paginated=True, timeout=timeout, with_ocr=with_ocr, ocr_language=ocr_language
            )
            result = converter.convert(path, raises_on_error=False, page_range=page_range)
        except Exception as error:
            logger.warning(
                "PDF conversion failed for pages %d-%d (ocr=%s): %s",
                *page_range,
                with_ocr,
                type(error).__name__,
            )
            return None
        if result.status not in (
            docling.ConversionStatus.SUCCESS,
            docling.ConversionStatus.PARTIAL_SUCCESS,
        ):
            logger.warning(
                "PDF conversion failed for pages %d-%d (ocr=%s): status %s",
                *page_range,
                with_ocr,
                result.status.value,
            )
            return None
        return _PageConversion(
            texts=_export_pages(result.document, page_range),
            partial=result.status == docling.ConversionStatus.PARTIAL_SUCCESS,
        )

    def _convert_whole(
        self,
        path: Path,
        source_format: SourceFormat,
        progress: ProgressCallback,
        base_url: str,
        ocr_language: str,
    ) -> ConversionOutput:
        docling = _docling()
        progress(0, 1)
        is_image = source_format in IMAGE_FORMATS
        failed = SourceError("conversion_failed", "This file could not be converted.")
        source: Any = path
        has_content_images = False
        if source_format == SourceFormat.HTML:
            converter = self._html_converter(base_url)
            data = _normalize_html(path.read_bytes())
            has_content_images = _HTML_IMAGE_TAG.search(data) is not None
            # From a file path, Docling resolves links like "/about/" against the temp folder
            # and fails the whole page; a stream uses `base_url`, or keeps links as written.
            source = docling.DocumentStream(name=path.name, stream=io.BytesIO(data))
        try:
            if source_format != SourceFormat.HTML:
                converter = self._converter(
                    paginated=is_image,
                    timeout=MAX_TIMEOUT,
                    with_ocr=is_image,
                    ocr_language=ocr_language,
                )
            result = converter.convert(source, raises_on_error=False)
        except Exception as error:
            logger.warning("Conversion failed: %s", type(error).__name__)
            raise failed from error
        if result.status not in (
            docling.ConversionStatus.SUCCESS,
            docling.ConversionStatus.PARTIAL_SUCCESS,
        ):
            logger.warning(
                "Conversion failed: status %s, %d error(s)", result.status.value, len(result.errors)
            )
            raise failed

        document = result.document
        sheet_names = _sheet_names(path) if source_format == SourceFormat.XLSX else []
        if document.pages:
            numbers = sorted(document.pages)
            use_sheets = len(sheet_names) == len(numbers)
            units = [
                ConvertedUnit(
                    page=number,
                    markdown=document.export_to_markdown(page_no=number),
                    sheet=sheet_names[index] if use_sheets else "",
                )
                for index, number in enumerate(numbers)
            ]
        else:
            units = [ConvertedUnit(page=None, markdown=document.export_to_markdown())]
        missing = [unit.page for unit in units if not _text_tokens(unit.markdown)]
        units = [unit for unit in units if _text_tokens(unit.markdown)]
        if not units:
            if source_format == SourceFormat.HTML:
                if has_content_images:
                    raise SourceError(
                        "no_text",
                        "No readable text was found on this page; its content appears to be "
                        "images. Text in images on web pages is not read. Add the image "
                        "addresses directly as sources, or upload copies of the images.",
                    )
                raise SourceError(
                    "no_text",
                    "The page has little or no readable text. It may need JavaScript to show "
                    "its content; upload the document or save the page as PDF instead.",
                )
            raise SourceError("no_text", "No readable text was found in this file.")
        limitations: list[str] = []
        if result.status == docling.ConversionStatus.PARTIAL_SUCCESS:
            limitations.append("Part of this file could not be converted.")
        if source_format == SourceFormat.HTML and _html_has_sparse_body(
            "\n\n".join(unit.markdown for unit in units)
        ):
            limitations.append(
                "This page contains many headings but very little body text. "
                "FAQ answers or other content may be missing. Check the converted text "
                "against the original page, or upload a complete document."
            )
        if is_image and missing:
            limitations.append("Some image pages contained no readable text.")
        progress(1, 1)
        return ConversionOutput(units=units, limitations=limitations)


def _export_pages(document: Any, page_range: tuple[int, int]) -> dict[int, str]:
    start, end = page_range
    numbers = sorted(document.pages)
    # Keep original page numbers even if Docling numbered the range from 1.
    offset = 0 if all(start <= number <= end for number in numbers) else start - 1
    return {number + offset: document.export_to_markdown(page_no=number) for number in numbers}


def _sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True)
    except Exception:
        return []
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()
